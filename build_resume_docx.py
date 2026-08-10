#!/usr/bin/env python3
"""Build an ATS-friendly .docx resume.

Deliberately plain: single column, no tables, no text boxes, no headers/footers,
standard fonts. Applicant tracking systems parse this cleanly; fancy templates
are where parsing breaks.

Usage: python3 build_resume_docx.py [output.docx]
"""

import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

NAME = "Chinedu Ezeofor"
TITLE = "Service Desk Analyst II"
CONTACT = "New York, NY  |  chinedu@interactxp.com  |  linkedin.com/in/c-ezeofor90"

SUMMARY = (
    "Service desk analyst working in IT support since 2021, currently covering the weekend shift "
    "for a healthcare client with 1,000+ users. Most of my day is Active Directory, Epic, MyChart "
    "and Microsoft 365 - creating accounts, clearing blocks, verifying who I'm actually talking to "
    "before I change anything, and talking users through software changes they didn't ask for. "
    "Outside of work I build automations in Make.com and n8n and run a small digital agency in NYC. "
    "Studying for Security+."
)

EXPERIENCE = [
    {
        "role": "Service Desk Analyst II",
        "org": "Matrix Global Services",
        "dates": "April 2023 - Present",
        "place": "Remote",
        "bullets": [
            "Create user accounts in on-prem Active Directory (Windows Server AD DS) for a hybrid "
            "Microsoft 365 environment, set the username and email format, and add users to the "
            "groups and shared inboxes I have rights to. Anything outside that list goes to the "
            "mail team as a ticket.",
            "Reset passwords and clear account blocks across Active Directory, Epic, MyChart and "
            "Healthstream. Every request goes through caller verification and a Duo push before I "
            "change anything.",
            "Handle MyChart proxy access carefully - confirm proxy rights, or get the account "
            "owner's consent on the call. Callers who can't verify get routed to Medical Records.",
            "Own status calls start to finish: note the callback on the original ticket, reopen it "
            "and raise the priority when the problem is still happening, nudge whoever owns it, "
            "then close the status ticket with what the user told me.",
            "Walk users through migrations they're resistant to. Moving people off the old intranet "
            "webmail onto Outlook in Microsoft 365 mostly meant showing them they could save it as "
            "a PWA and still have an icon to click.",
            "Process name changes and transfers, and reinstate returning employees once HR or their "
            "supervisor approves the form.",
            "Escalate what's locked down - printer installs behind UAC, anything past a small tweak "
            "- with the diagnosis and steps already written up for site support.",
        ],
        "env": "Environment: Citrix Workspace, internal RDP, Zoho Assist, FreshService, Active "
               "Directory, Microsoft 365, Epic, MyChart, Healthstream, Duo",
    },
    {
        "role": "IT Support Specialist (Google x Multiverse Apprenticeship)",
        "org": "Google",
        "dates": "August 2021 - January 2023",
        "place": "New York, NY",
        "bullets": [
            "Supported Google employees over chat, call and ticket across Windows, macOS and mobile "
            "devices.",
            "Worked with subject matter experts to provision security keys.",
            "Finished the apprenticeship with the Google IT Support Professional Certificate.",
        ],
    },
    {
        "role": "Founder",
        "org": "Interactive Xperience Agency",
        "dates": "",
        "place": "New York, NY",
        "bullets": [
            "Run a small digital agency for NYC small businesses: local SEO, lead generation, "
            "reputation management.",
            "Built a Make.com workflow that takes Facebook group join requests, verifies the email "
            "address, and passes everyone who clears into a clean marketing list.",
            "Connected the OpenAI API to a spreadsheet so copy can be generated and questions asked "
            "against the data already in the sheet, without leaving it.",
        ],
    },
    {
        "role": "Clerical Support",
        "org": "New York City Law Department",
        "dates": "July 2017 - August 2021",
        "place": "Brooklyn, NY",
        "bullets": [
            "Supported legal staff in a document-heavy office running iManage.",
            "Answered questions from clients and the public, and kept filings and records organized.",
        ],
    },
]

SKILLS = [
    ("Identity & access", "Active Directory (Windows Server AD DS), Microsoft 365, Duo MFA, "
                          "account provisioning, group and shared mailbox access"),
    ("Service desk", "FreshService (ITSM), Zoho Assist, Citrix Workspace, remote desktop, "
                     "ticket triage and escalation"),
    ("Healthcare systems", "Epic, MyChart, Healthstream"),
    ("Platforms", "Windows, macOS, iOS, Android"),
    ("Networking", "TCP/IP, DNS, DHCP, VPN, LAN/Wi-Fi, ping / ipconfig / nslookup / tracert"),
    ("Automation & AI", "Make.com, n8n, OpenAI API, no-code databases, prompt writing"),
]

CERTS = [
    "Google IT Support Professional Certificate - Google x Multiverse Apprenticeship",
    "CompTIA Security+ - in progress",
]

PROJECTS = [
    "Built and maintain a custom desktop PC: parts selection, assembly, OS install, hardware "
    "troubleshooting.",
    "Run and troubleshoot my own home network.",
]

ACCENT = RGBColor(0x1A, 0x3D, 0x6D)


def setup(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.05


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = ACCENT
    # A simple bottom border reads fine in ATS parsers (it is paragraph
    # formatting, not a table or graphic).
    pr = p._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A3D6D")
    borders.append(bottom)
    pr.append(borders)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.22)
    return p


def build(path):
    doc = Document()
    setup(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(NAME)
    r.bold = True
    r.font.size = Pt(20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    r.font.size = Pt(11.5)
    r.font.color.rgb = ACCENT
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run(CONTACT).font.size = Pt(9.5)

    heading(doc, "Summary")
    doc.add_paragraph(SUMMARY).paragraph_format.space_after = Pt(2)

    heading(doc, "Experience")
    for job in EXPERIENCE:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        r = p.add_run(f"{job['role']} - {job['org']}")
        r.bold = True
        r.font.size = Pt(10.5)

        meta = " | ".join(x for x in (job["dates"], job["place"]) if x)
        if meta:
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_after = Pt(2)
            r2 = p2.add_run(meta)
            r2.italic = True
            r2.font.size = Pt(9)

        for b in job["bullets"]:
            bullet(doc, b)

        if job.get("env"):
            p3 = doc.add_paragraph()
            p3.paragraph_format.space_before = Pt(2)
            r3 = p3.add_run(job["env"])
            r3.font.size = Pt(9)
            r3.italic = True

    heading(doc, "Technical Skills")
    for label, items in SKILLS:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"{label}: ")
        r.bold = True
        p.add_run(items)

    heading(doc, "Certifications")
    for c in CERTS:
        bullet(doc, c)

    heading(doc, "Projects")
    for pr_ in PROJECTS:
        bullet(doc, pr_)

    doc.save(path)
    return path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "Chinedu_Ezeofor_Resume.docx"
    print("wrote", build(out))
