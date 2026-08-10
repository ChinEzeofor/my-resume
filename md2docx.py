#!/usr/bin/env python3
"""Convert a resume markdown file to an ATS-safe .docx.

resume.md and resume-automation.md are the single source of truth; this script
renders them. Output is deliberately plain - single column, no tables, no text
boxes, no headers/footers - because that is what applicant tracking systems
parse reliably.

Usage:
    python3 md2docx.py resume.md Chinedu_Ezeofor_Resume.docx
"""

import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x1A, 0x3D, 0x6D)


def _rule(p):
    """Bottom border on a section heading. Paragraph formatting, not a table."""
    pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A3D6D")
    borders.append(bottom)
    pr.append(borders)


def _clean(text):
    """Strip markdown emphasis and swap typographic characters ATS parsers mangle."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    return text.replace("—", "-").replace("–", "-").replace("·", "|").replace("×", "x").strip()


def build(md_path, out_path):
    lines = open(md_path).read().split("\n")
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.5)
        section.left_margin = section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.05

    buf = []

    def flush():
        if buf:
            p = doc.add_paragraph(_clean(" ".join(buf)))
            p.paragraph_format.space_after = Pt(3)
            buf.clear()

    i = 0
    while i < len(lines):
        s = lines[i].strip()

        if s.startswith("> "):          # internal note, never rendered
            i += 1
            continue

        if s.startswith("# "):          # name
            flush()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(_clean(s[2:]))
            r.bold = True
            r.font.size = Pt(20)

        elif s.startswith("## "):       # section heading
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(_clean(s[3:]).upper())
            r.bold = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = ACCENT
            _rule(p)

        elif s.startswith("### "):      # job title
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            r = p.add_run(_clean(s[4:]))
            r.bold = True
            r.font.size = Pt(10.5)

        elif s.startswith("**") and s.endswith("**") and len(s) > 4 and not s.startswith("**Environment"):
            flush()                     # headline under the name
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(_clean(s))
            r.bold = True
            r.font.size = Pt(11.5)
            r.font.color.rgb = ACCENT

        elif s.startswith("*") and s.endswith("*") and not s.startswith("**"):
            flush()                     # dates / location line
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(_clean(s))
            r.italic = True
            r.font.size = Pt(9)

        elif s.startswith("- "):        # bullet, joining wrapped continuation lines
            flush()
            parts = [s[2:]]
            while i + 1 < len(lines) and lines[i + 1].startswith("  ") and not lines[i + 1].strip().startswith("-"):
                i += 1
                parts.append(lines[i].strip())
            p = doc.add_paragraph(_clean(" ".join(parts)), style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.22)

        elif not s:
            flush()
        else:
            buf.append(s)
        i += 1

    flush()
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "resume.md"
    dst = sys.argv[2] if len(sys.argv) > 2 else "Chinedu_Ezeofor_Resume.docx"
    print("wrote", build(src, dst))
